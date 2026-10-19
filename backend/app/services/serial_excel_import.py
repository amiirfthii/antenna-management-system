#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Import serial-number antennas from column-oriented Excel.

Excel layout (per column):
  row 1: entry date (Jalali)
  row 2: vendor
  row 3: antenna type
  row 4+: serial numbers

Item code is resolved from itemCodeAssignments.
Stock in noSerialAntennas is decremented:
  - 5X / 8X / TDD → first استان; if empty, fallback to مستهلک
  - TDD ignores exact antenna type/vendor and deducts from any TDD stock pool
  - otherwise → category مستهلک
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

import jdatetime
import pandas as pd
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.antenna_type_bucket import classify_antenna_x_bucket
from app.core.jalali_dates import normalize_jalali_date_string
from app.db import models
from app.schemas.province_antenna import CATEGORY_DEPRECIATED, CATEGORY_PROVINCE
from app.services.no_serial_antenna import consume_no_serial_unit, is_tdd_label
from app.services.serial_uniqueness import find_repair_by_serial, normalize_serial, serial_key


def _norm(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _norm_key(value: Any) -> str:
    return _norm(value).lower()


def _parse_entry_date(value: Any) -> Optional[str]:
    raw = _norm(value)
    if not raw:
        return None
    # Excel may store dates as datetime
    if hasattr(value, "strftime") and not isinstance(value, str):
        try:
            # unlikely gregorian in this workflow; keep string path
            raw = str(value)
        except Exception:
            pass
    normalized = normalize_jalali_date_string(raw)
    return normalized if normalized else raw


def resolve_stock_category(vendor: str, antenna_type: str) -> str:
    """5X / 8X / TDD → استان (با فال‌بک مستهلک)؛ بقیه → مستهلک."""
    blob = f"{vendor} {antenna_type}".upper()
    if "TDD" in blob:
        return CATEGORY_PROVINCE
    bucket = classify_antenna_x_bucket(antenna_type)
    if bucket in ("5X", "8X"):
        return CATEGORY_PROVINCE
    return CATEGORY_DEPRECIATED


def uses_province_then_depreciated(vendor: str, antenna_type: str) -> bool:
    return resolve_stock_category(vendor, antenna_type) == CATEGORY_PROVINCE


def _lookup_item_code(
    db: Session,
    vendor: str,
    antenna_type: str,
    vendor_map: Dict[str, str],
    type_map: Dict[str, str],
) -> Optional[str]:
    v_key = _norm_key(vendor)
    t_key = _norm_key(antenna_type)
    for row in db.query(models.ItemCodeAssignment).all():
        m = _norm_key(row.manufacturer)
        a = _norm_key(row.antenna_type)
        if m == v_key and a == t_key:
            code = _norm(row.item_code)
            if code:
                return code
    # also try canonical names
    canon_v = vendor_map.get(v_key, vendor)
    canon_t = type_map.get(t_key, antenna_type)
    if canon_v != vendor or canon_t != antenna_type:
        return _lookup_item_code(db, canon_v, canon_t, {}, {})
    return None


def _parse_columns(content: bytes) -> List[Dict[str, Any]]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    max_col = ws.max_column or 0
    max_row = ws.max_row or 0
    columns: List[Dict[str, Any]] = []

    for col in range(1, max_col + 1):
        entry_date = _parse_entry_date(ws.cell(1, col).value)
        vendor = _norm(ws.cell(2, col).value)
        antenna_type = _norm(ws.cell(3, col).value)
        if not entry_date and not vendor and not antenna_type:
            # skip fully empty header columns
            serials_probe = [
                _norm(ws.cell(r, col).value)
                for r in range(4, min(max_row, 6) + 1)
            ]
            if not any(serials_probe):
                continue

        serials: List[Tuple[int, str]] = []
        for row in range(4, max_row + 1):
            serial = _norm(ws.cell(row, col).value)
            if serial:
                serials.append((row, serial))

        columns.append(
            {
                "col": col,
                "entry_date": entry_date,
                "vendor": vendor,
                "antenna_type": antenna_type,
                "serials": serials,
            }
        )
    return columns


def import_serials_from_excel(db: Session, content: bytes) -> Dict[str, Any]:
    try:
        columns = _parse_columns(content)
    except Exception as exc:
        return {
            "total": 0,
            "imported": 0,
            "failed": 0,
            "errors": [f"خواندن فایل اکسل ناموفق بود: {exc}"],
            "imported_rows": [],
        }

    vendor_map = {
        _norm_key(r.company_constructor): _norm(r.company_constructor)
        for r in db.query(models.AntennaConstructor).all()
        if _norm(r.company_constructor)
    }
    type_map = {
        _norm_key(r.antenna_type): _norm(r.antenna_type)
        for r in db.query(models.AntennaType).all()
        if _norm(r.antenna_type)
    }

    total = imported = failed = 0
    errors: List[str] = []
    imported_rows: List[Dict[str, Any]] = []
    seen_in_file: set[str] = set()

    for col_info in columns:
        col = col_info["col"]
        entry_date = col_info["entry_date"] or jdatetime.datetime.now().strftime("%Y/%m/%d")
        raw_vendor = col_info["vendor"]
        raw_type = col_info["antenna_type"]

        if not raw_vendor or not raw_type:
            for row_no, serial in col_info["serials"]:
                total += 1
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no} (سریال {serial}): وندور یا نوع آنتن خالی است"
                )
            if not col_info["serials"]:
                failed += 1
                errors.append(f"ستون {col}: وندور یا نوع آنتن خالی است")
            continue

        vendor = vendor_map.get(_norm_key(raw_vendor), raw_vendor)
        antenna_type = type_map.get(_norm_key(raw_type), raw_type)

        if _norm_key(raw_vendor) not in vendor_map:
            for row_no, serial in col_info["serials"]:
                total += 1
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no} (سریال {serial}): وندور «{raw_vendor}» در دیتابیس نیست"
                )
            continue

        if _norm_key(raw_type) not in type_map:
            for row_no, serial in col_info["serials"]:
                total += 1
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no} (سریال {serial}): نوع آنتن «{raw_type}» در دیتابیس نیست"
                )
            continue

        item_code = _lookup_item_code(db, vendor, antenna_type, vendor_map, type_map)
        if not item_code:
            for row_no, serial in col_info["serials"]:
                total += 1
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no} (سریال {serial}): "
                    f"آیتم‌کد برای «{vendor} / {antenna_type}» یافت نشد"
                )
            continue

        category = resolve_stock_category(vendor, antenna_type)

        for row_no, serial in col_info["serials"]:
            total += 1
            serial = normalize_serial(serial)
            if not serial:
                failed += 1
                errors.append(f"ستون {col} ردیف {row_no}: سریال خالی است")
                continue

            serial_lookup = serial_key(serial)
            if serial_lookup in seen_in_file:
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no}: سریال تکراری در فایل ({serial})"
                )
                continue
            seen_in_file.add(serial_lookup)

            existing = find_repair_by_serial(db, serial)
            if existing:
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no}: سریال {serial} قبلاً ثبت شده "
                    f"(تعمیر {existing.repair_no})"
                )
                continue

            try:
                fallback = (
                    [CATEGORY_DEPRECIATED]
                    if uses_province_then_depreciated(vendor, antenna_type)
                    else None
                )
                used_category = consume_no_serial_unit(
                    db,
                    vendor=vendor,
                    antenna_type=antenna_type,
                    item_code=item_code,
                    category=category,
                    allow_item_code_fallback=True,
                    fallback_categories=fallback,
                    tdd_pool=is_tdd_label(vendor, antenna_type),
                )
            except ValueError as exc:
                failed += 1
                errors.append(
                    f"ستون {col} ردیف {row_no} (سریال {serial}): {exc}"
                )
                continue

            repair = models.Repair(
                antenna_serial_no=serial,
                antenna_type=antenna_type,
                constructor=vendor,
                antenna_entry_type="ورودی_تعمیر",
                entry_date=entry_date,
                item_code=item_code,
                antenna_status="نامعلوم",
                no_serial_category=used_category,
            )
            db.add(repair)
            db.flush()
            imported += 1
            imported_rows.append(
                {
                    "repair_no": repair.repair_no,
                    "serial": serial,
                    "vendor": vendor,
                    "antenna_type": antenna_type,
                    "item_code": item_code,
                    "entry_date": entry_date,
                    "category": used_category,
                }
            )

    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        return {
            "total": total,
            "imported": 0,
            "failed": total,
            "errors": [f"خطا در ذخیره نهایی: {exc}"],
            "imported_rows": [],
        }

    return {
        "total": total,
        "imported": imported,
        "failed": failed,
        "errors": errors[:200],
        "imported_rows": imported_rows[:200],
    }
