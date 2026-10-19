#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Parse consumable inventory snapshot Excel (موجودی انبار *.xlsx)."""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import List

import openpyxl


@dataclass
class ConsumableSnapshotRow:
    name: str
    initial_qty: int
    current_qty: int
    consumed_qty: int


def _norm_name(value: str) -> str:
    s = unicodedata.normalize("NFKC", (value or "").strip().lower())
    s = s.replace("ك", "ک").replace("ي", "ی")
    s = s.replace("فیکساتور", "فیکاساتور")
    s = re.sub(r"\s+", " ", s)
    return s


def find_inventory_xlsx(search_dir: Path | None = None) -> Path:
    if search_dir is not None:
        roots = [search_dir]
    else:
        here = Path(__file__).resolve()
        roots = [here.parents[5], here.parents[4], here.parents[3]]
    for root in roots:
        if not root.exists():
            continue
        matches = sorted(root.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
        for p in matches:
            if "04.29" in p.name or "موجودی" in p.name:
                return p
    raise FileNotFoundError(f"No inventory xlsx found under {roots}")


def parse_consumable_snapshot_xlsx(path: Path) -> List[ConsumableSnapshotRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows: List[ConsumableSnapshotRow] = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name or not str(name).strip():
            continue
        initial_qty = int(ws.cell(r, 2).value or 0)
        current_qty = int(ws.cell(r, 3).value or 0)
        consumed_qty = int(ws.cell(r, 4).value or 0)
        if consumed_qty <= 0 and initial_qty > 0:
            consumed_qty = max(0, initial_qty - current_qty)
        if current_qty <= 0 and initial_qty > 0:
            current_qty = max(0, initial_qty - consumed_qty)
        rows.append(
            ConsumableSnapshotRow(
                name=str(name).strip(),
                initial_qty=initial_qty,
                current_qty=current_qty,
                consumed_qty=consumed_qty,
            )
        )
    return rows


def match_part(db_name: str, excel_name: str) -> bool:
    a = _norm_name(db_name)
    b = _norm_name(excel_name)
    return a == b or a in b or b in a


def find_best_match(items, excel_name: str):
    """Return best ArsenalTable row for an Excel part name."""
    norm_excel = _norm_name(excel_name)
    exact = [i for i in items if _norm_name(i.fragment_name or "") == norm_excel]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        return exact[0]
    partial = [i for i in items if match_part(i.fragment_name or "", excel_name)]
    if len(partial) == 1:
        return partial[0]
    return None


__all__ = [
    "ConsumableSnapshotRow",
    "find_inventory_xlsx",
    "parse_consumable_snapshot_xlsx",
    "match_part",
    "find_best_match",
    "_norm_name",
]
