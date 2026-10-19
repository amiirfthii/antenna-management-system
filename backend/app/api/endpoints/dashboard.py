#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional
import json

import jdatetime
from fastapi import APIRouter, Body, Depends, HTTPException, Query, Response, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_current_admin_user
from app.core.antenna_type_bucket import antenna_x_bucket_sort_key, classify_antenna_x_bucket
from app.core.arsenal import ARSENAL_EQUIPMENT
from app.core.jalali_dates import apply_jalali_date_range_filter
from app.core.profile_photos import get_photo_url
from app.db import models
from app.db.database import get_db
from app.services.pmo_display_overrides import (
    clear_pmo_display_overrides,
)
from app.services.pmo_snapshot import (
    bars_from_snapshot,
    center_inventory_from_snapshot,
    load_snapshot,
    save_snapshot,
    seed_if_empty,
    snapshot_card_totals,
    snapshot_has_rows,
)
from app.services.serial_uniqueness import serial_key as repair_serial_key

from app.schemas.province_antenna import CATEGORY_DEPRECIATED
from app.services.no_serial_antenna import remaining_of

router = APIRouter()

STATUS_LABELS: Dict[str, str] = {
    "نامعلوم": "نامعلوم",
    "تایید": "تایید",
    "فالتی_در_انتظار_تعمیر": "فالتی در انتظار بعدی",
    "در_انتظار_تایید": "در انتظار تایید",
    "خروج": "خروج",
    "خروج_تایید_شده": "خروج تایید شده",
    "خروج_تایید_نشده": "خروج تایید نشده",
    "فالتی_سالویج": "فالتی سالویج",
    "فالتی_در_انتظار_تایید": "فالتی در انتظار تایید",
}

ITEM_CODE_STATUS_ORDER = [
    "تایید",
    "خروج",
    "نامعلوم",
    "فالتی_در_انتظار_تعمیر",
    "فالتی_سالویج",
]

ITEM_CODE_STATUS_COLORS: Dict[str, str] = {
    "تایید": "#22c55e",
    "خروج": "#3b82f6",
    "نامعلوم": "#f59e0b",
    "فالتی_در_انتظار_تعمیر": "#ef4444",
    "فالتی_سالویج": "#a855f7",
}

EXIT_STATUSES = frozenset(
    {
        "خروج",
        "خروج_تایید_شده",
        "خروج_تایید_نشده",
    }
)


def _is_exit_status(status: str) -> bool:
    return status in EXIT_STATUSES or status.startswith("خروج")


def _item_code_label_map(db: Session) -> dict[str, str]:
    """Map item codes to 'vendor - antenna type' for chart labels."""
    labels: dict[str, str] = {}
    if not hasattr(models, "ItemCodeAssignment"):
        return labels
    for row in (
        db.query(
            models.ItemCodeAssignment.item_code,
            models.ItemCodeAssignment.manufacturer,
            models.ItemCodeAssignment.antenna_type,
        )
        .all()
    ):
        code = (row.item_code or "").strip()
        if not code or code in labels:
            continue
        vendor = (row.manufacturer or "").strip()
        antenna_type = (row.antenna_type or "").strip()
        if vendor and antenna_type:
            labels[code] = f"{vendor} - {antenna_type}"
        elif vendor:
            labels[code] = vendor
        elif antenna_type:
            labels[code] = antenna_type
    return labels


def _build_item_code_bars(db: Session) -> List[Dict[str, Any]]:
    """Per item code: approved, exited, and province incoming."""
    grouped: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))

    for repair in (
        db.query(models.Repair.item_code, models.Repair.antenna_status)
        .filter(
            models.Repair.item_code.isnot(None),
            models.Repair.item_code != "",
        )
        .all()
    ):
        code = (repair.item_code or "").strip()
        if not code:
            continue
        status = repair.antenna_status or "نامعلوم"
        grouped[code][status] += 1

    province_by_code: dict[str, int] = defaultdict(int)
    if hasattr(models, "NoSerialAntenna"):
        for row in (
            db.query(
                models.NoSerialAntenna.item_code,
                models.NoSerialAntenna.quantity,
                models.NoSerialAntenna.remaining_quantity,
            )
            .filter(models.NoSerialAntenna.category == "استان")
            .all()
        ):
            code = (row.item_code or "").strip()
            if code:
                province_by_code[code] += remaining_of(row)

    assignment_codes: set[str] = set()
    if hasattr(models, "ItemCodeAssignment"):
        for row in db.query(models.ItemCodeAssignment.item_code).distinct().all():
            code = (row.item_code or "").strip()
            if code:
                assignment_codes.add(code)

    all_codes = set(grouped.keys()) | set(province_by_code.keys()) | assignment_codes
    label_map = _item_code_label_map(db)
    bars: List[Dict[str, Any]] = []
    for code in all_codes:
        status_counts = grouped.get(code, {})
        total = sum(status_counts.values())
        approved = int(status_counts.get("تایید", 0))
        exited = sum(count for status, count in status_counts.items() if _is_exit_status(status))
        in_center = max(0, total - exited)
        province_in = province_by_code.get(code, 0)
        bars.append(
            {
                "code": code,
                "label": label_map.get(code, code),
                "approved": approved,
                "in_center": in_center,
                "exited": exited,
                "province_in": province_in,
                "total": approved + exited + province_in,
            }
        )

    bars.sort(
        key=lambda row: (
            -int(row["approved"] + row["exited"] + row["province_in"]),
            row["code"],
        )
    )
    return bars


def _item_code_to_antenna_type_map(db: Session) -> dict[str, str]:
    mapping: dict[str, str] = {}
    if not hasattr(models, "ItemCodeAssignment"):
        return mapping
    for row in db.query(
        models.ItemCodeAssignment.item_code,
        models.ItemCodeAssignment.antenna_type,
    ).all():
        code = (row.item_code or "").strip().lower()
        antenna_type = (row.antenna_type or "").strip()
        if code and antenna_type and code not in mapping:
            mapping[code] = antenna_type
    return mapping


def _resolve_antenna_type(
    antenna_type: str | None,
    item_code: str | None,
    item_type_map: dict[str, str],
) -> str:
    resolved = (antenna_type or "").strip()
    if resolved:
        return resolved
    code = (item_code or "").strip().lower()
    if code:
        return item_type_map.get(code, "")
    return ""


def _build_antenna_type_x_bars(db: Session) -> List[Dict[str, Any]]:
    """Aggregate approved / exit / province-in by antenna X class (2X, 3X, …)."""
    grouped: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    item_type_map = _item_code_to_antenna_type_map(db)

    for repair in (
        db.query(
            models.Repair.item_code,
            models.Repair.antenna_type,
            models.Repair.antenna_status,
        )
        .all()
    ):
        antenna_type = _resolve_antenna_type(
            repair.antenna_type,
            repair.item_code,
            item_type_map,
        )
        bucket = classify_antenna_x_bucket(antenna_type)
        status = repair.antenna_status or "نامعلوم"
        grouped[bucket][status] += 1

    province_by_bucket: dict[str, int] = defaultdict(int)
    if hasattr(models, "NoSerialAntenna"):
        for row in (
            db.query(
                models.NoSerialAntenna.antenna_type,
                models.NoSerialAntenna.quantity,
                models.NoSerialAntenna.remaining_quantity,
            )
            .filter(models.NoSerialAntenna.category == "استان")
            .all()
        ):
            bucket = classify_antenna_x_bucket(row.antenna_type)
            province_by_bucket[bucket] += remaining_of(row)

    all_buckets = set(grouped.keys()) | set(province_by_bucket.keys())
    bars: List[Dict[str, Any]] = []
    for bucket in all_buckets:
        status_counts = grouped.get(bucket, {})
        total = sum(status_counts.values())
        approved = int(status_counts.get("تایید", 0))
        exited = sum(count for status, count in status_counts.items() if _is_exit_status(status))
        in_center = max(0, total - exited)
        province_in = province_by_bucket.get(bucket, 0)
        if approved + exited + province_in <= 0:
            continue
        bars.append(
            {
                "code": bucket,
                "label": bucket,
                "approved": approved,
                "in_center": in_center,
                "exited": exited,
                "province_in": province_in,
                "total": approved + exited + province_in,
            }
        )

    bars.sort(
        key=lambda row: (
            antenna_x_bucket_sort_key(row["code"]),
            -int(row["approved"] + row["exited"] + row["province_in"]),
        )
    )
    return bars


def _build_center_inventory_bars(db: Session) -> List[Dict[str, Any]]:
    """Center inventory by antenna X class, split into province vs depreciated.

    Counts both serial (repairs, excluding خروج) and non-serial (NoSerialAntenna)
    antennas, grouped by their source category (استان / مستهلک).
    """
    province_by_bucket: dict[str, int] = defaultdict(int)
    depreciated_by_bucket: dict[str, int] = defaultdict(int)
    item_type_map = _item_code_to_antenna_type_map(db)

    # non-serial stock (quantities)
    if hasattr(models, "NoSerialAntenna"):
        for row in db.query(
            models.NoSerialAntenna.antenna_type,
            models.NoSerialAntenna.quantity,
            models.NoSerialAntenna.remaining_quantity,
            models.NoSerialAntenna.category,
        ).all():
            bucket = classify_antenna_x_bucket(row.antenna_type)
            qty = remaining_of(row)
            category = (row.category or "").strip()
            if category == "مستهلک":
                depreciated_by_bucket[bucket] += qty
            else:
                province_by_bucket[bucket] += qty

    # serial antennas currently in center (status != خروج) with a stored source category
    if hasattr(models.Repair, "no_serial_category"):
        for repair in db.query(
            models.Repair.item_code,
            models.Repair.antenna_type,
            models.Repair.antenna_status,
            models.Repair.no_serial_category,
        ).all():
            category = (repair.no_serial_category or "").strip()
            if category not in ("استان", "مستهلک"):
                continue
            status = repair.antenna_status or "نامعلوم"
            if _is_exit_status(status):
                continue
            antenna_type = _resolve_antenna_type(
                repair.antenna_type, repair.item_code, item_type_map
            )
            bucket = classify_antenna_x_bucket(antenna_type)
            if category == "مستهلک":
                depreciated_by_bucket[bucket] += 1
            else:
                province_by_bucket[bucket] += 1

    all_buckets = set(province_by_bucket.keys()) | set(depreciated_by_bucket.keys())
    bars: List[Dict[str, Any]] = []
    for bucket in all_buckets:
        province_in = int(province_by_bucket.get(bucket, 0))
        depreciated_in = int(depreciated_by_bucket.get(bucket, 0))
        if province_in + depreciated_in <= 0:
            continue
        bars.append(
            {
                "code": bucket,
                "label": bucket,
                "province_in": province_in,
                "depreciated_in": depreciated_in,
                "total": province_in + depreciated_in,
            }
        )

    bars.sort(
        key=lambda row: (
            antenna_x_bucket_sort_key(row["code"]),
            -int(row["province_in"] + row["depreciated_in"]),
        )
    )
    return bars


def _build_top_repairers(db: Session) -> List[Dict[str, Any]]:
    """Per active user: photo + their most-repaired item (by distinct antennas)."""
    if not hasattr(models, "RepairItemReport"):
        return []

    users = {
        u.personal_id: u
        for u in db.query(models.User).filter(models.User.is_active.is_(True)).all()
    }
    if not users:
        return []

    seen: dict[int, set] = defaultdict(set)
    counts: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for row in db.query(
        models.RepairItemReport.personal_id,
        models.RepairItemReport.item_name,
        models.RepairItemReport.antenna_serial_no,
    ).all():
        pid = row.personal_id
        if not pid or pid not in users:
            continue
        item = (row.item_name or "").strip() or "نامشخص"
        serial = (row.antenna_serial_no or "").strip()
        key = (serial, item)
        if key in seen[pid]:
            continue
        seen[pid].add(key)
        counts[pid][item] += 1

    result: List[Dict[str, Any]] = []
    for pid, user in users.items():
        cmap = counts.get(pid)
        if not cmap:
            continue
        top_item, top_count = max(cmap.items(), key=lambda kv: kv[1])
        total = sum(cmap.values())
        items_sorted = sorted(cmap.items(), key=lambda kv: -kv[1])[:5]
        first = (user.user_first_name or "").strip()
        last = (user.user_last_name or "").strip()
        result.append(
            {
                "personal_id": pid,
                "name": (first + " " + last).strip() or "کاربر",
                "photo_url": get_photo_url(first, last),
                "top_item": top_item,
                "top_count": int(top_count),
                "total": int(total),
                "items": [
                    {"name": name, "count": int(cnt)}
                    for name, cnt in items_sorted
                ],
            }
        )

    result.sort(key=lambda r: -r["top_count"])
    return result


def _build_vendor_top_units(db: Session) -> Dict[str, Any]:
    """Display-only vendor inventory from Excel export (vendor_inventory.json)."""
    data_path = Path(__file__).resolve().parents[2] / "data" / "vendor_inventory.json"
    if not data_path.exists():
        return {"vendors": [], "by_vendor": {}}
    try:
        payload = json.loads(data_path.read_text(encoding="utf-8"))
        return {
            "source": payload.get("source"),
            "total_qty": payload.get("total_qty"),
            "vendor_totals": payload.get("vendor_totals") or {},
            "vendors": payload.get("vendors") or [],
            "by_vendor": payload.get("by_vendor") or {},
        }
    except (json.JSONDecodeError, OSError):
        return {"vendors": [], "by_vendor": {}}


def _build_antenna_targets_map(db: Session) -> Dict[str, int]:
    """Map of {bucket_code: employer target qty} for the target chart."""
    if not hasattr(models, "AntennaTarget"):
        return {}
    result: Dict[str, int] = {}
    for row in db.query(models.AntennaTarget).all():
        code = (row.bucket_code or "").strip()
        if code:
            result[code] = int(row.target_qty or 0)
    return result


def _status_label(status: str) -> str:
    return STATUS_LABELS.get(status, status.replace("_", " "))


def _build_item_code_chart(status_counts: dict[str, int]) -> List[Dict[str, Any]]:
    chart: List[Dict[str, Any]] = []
    seen: set[str] = set()

    for status in ITEM_CODE_STATUS_ORDER:
        count = status_counts.get(status, 0)
        if count <= 0:
            continue
        seen.add(status)
        chart.append(
            {
                "status": status,
                "name": _status_label(status),
                "value": count,
                "color": ITEM_CODE_STATUS_COLORS.get(status, "#94a3b8"),
            }
        )

    for status, count in status_counts.items():
        if status in seen or count <= 0:
            continue
        chart.append(
            {
                "status": status,
                "name": _status_label(status),
                "value": count,
                "color": ITEM_CODE_STATUS_COLORS.get(status, "#94a3b8"),
            }
        )

    return chart


def _mode_label(counts: dict[str, int]) -> str:
    if not counts:
        return "-"
    return max(counts, key=counts.get)


def _build_category_inventory(db: Session, category: str) -> Dict[str, Any]:
    """موجودی مرکز برای یک دسته (مستهلک/استان): تفکیک وندور → نوع آنتن از NoSerialAntenna.

    خروجی هم‌شکل فایل قدیمی vendor_inventory.json است تا پنل داشبورد بدون تغییر ساختار
    آن را نمایش دهد.
    """
    empty = {
        "source": "database",
        "category": category,
        "total_qty": 0,
        "vendors": [],
        "vendor_totals": {},
        "by_vendor": {},
    }
    if not hasattr(models, "NoSerialAntenna"):
        return empty

    by_vendor_type: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    vendor_totals: dict[str, int] = defaultdict(int)
    total = 0

    for row in (
        db.query(
            models.NoSerialAntenna.vendor,
            models.NoSerialAntenna.antenna_type,
            models.NoSerialAntenna.quantity,
            models.NoSerialAntenna.remaining_quantity,
        )
        .filter(models.NoSerialAntenna.category == category)
        .all()
    ):
        vendor = (row.vendor or "").strip() or "نامشخص"
        antenna_type = (row.antenna_type or "").strip() or "نامشخص"
        qty = remaining_of(row)
        if qty == 0:
            continue
        by_vendor_type[vendor][antenna_type] += qty
        vendor_totals[vendor] += qty
        total += qty

    vendors = sorted(vendor_totals.keys(), key=lambda v: (-vendor_totals[v], v))
    by_vendor = {
        vendor: [
            {
                "unit": antenna_type,
                "description": antenna_type,
                "type": antenna_type,
                "count": int(count),
            }
            for antenna_type, count in sorted(
                types.items(), key=lambda kv: (-kv[1], kv[0])
            )
        ]
        for vendor, types in by_vendor_type.items()
    }

    return {
        "source": "database",
        "category": category,
        "total_qty": int(total),
        "vendors": vendors,
        "vendor_totals": {k: int(v) for k, v in vendor_totals.items()},
        "by_vendor": by_vendor,
    }


def _build_depreciated_inventory(db: Session) -> Dict[str, Any]:
    """موجودی مستهلک مرکز (زنده از دیتابیس)."""
    return _build_category_inventory(db, CATEGORY_DEPRECIATED)


def _build_province_inventory(db: Session) -> Dict[str, Any]:
    """موجودی ورودی استان مرکز (زنده از دیتابیس)."""
    return _build_category_inventory(db, "استان")


# ---------------------------------------------------------------------------
# PMO — گزارش پروژه (زنده از دیتابیس)
# ---------------------------------------------------------------------------

# وندورهای «۸ ایکس ایرانی» (8x-irani):
# هر آنتن 8X که وندور آن در این مجموعه باشد در ردیف «8x-irani» شمرده می‌شود و
# بقیه‌ی 8X در ردیف «8x-pole». برای افزودن/حذف وندور ایرانی، فقط همین یک خط را
# ویرایش کنید؛ مثال: PMO_IRANI_VENDORS = frozenset({"Faravid", "Faraz", "Paya"})
PMO_IRANI_VENDORS: frozenset = frozenset()

PMO_ROW_ORDER: List[str] = [
    "x-pole",
    "v-pole",
    "2x-pole",
    "3x-pole",
    "4x-pole",
    "5x-pole",
    "8x-pole",
    "TDD",
    "8x-irani",
]

PMO_ROW_LABELS: Dict[str, str] = {
    "x-pole": "X-Pole",
    "v-pole": "V-Pole",
    "2x-pole": "2X-Pole",
    "3x-pole": "3X-Pole",
    "4x-pole": "4X-Pole",
    "5x-pole": "5X-Pole",
    "8x-pole": "8X-Pole",
    "TDD": "TDD",
    "8x-irani": "8X ایرانی",
}

PMO_SKIP_ROW = "__skip__"

_PMO_BUCKET_TO_ROW: Dict[str, str] = {
    "X-Pole": "x-pole",
    "V-Pole": "v-pole",
    "2X": "2x-pole",
    "3X": "3x-pole",
    "4X": "4x-pole",
    "5X": "5x-pole",
}


def _is_tdd(vendor: str, antenna_type: str) -> bool:
    return "tdd" in f"{vendor or ''} {antenna_type or ''}".lower()


def _pmo_row_for_auto(vendor: str, antenna_type: str) -> str | None:
    """تشخیص خودکار ردیف PMO (بدون نگاشت دستی)."""
    vendor = (vendor or "").strip()
    antenna_type = (antenna_type or "").strip()
    if _is_tdd(vendor, antenna_type):
        return "TDD"
    bucket = classify_antenna_x_bucket(antenna_type)
    if bucket == "8X":
        return "8x-irani" if vendor in PMO_IRANI_VENDORS else "8x-pole"
    return _PMO_BUCKET_TO_ROW.get(bucket)


def _pmo_row_for(
    vendor: str,
    antenna_type: str,
    mapping_map: Dict[str, str] | None = None,
) -> str | None:
    """نگاشت (وندور، نوع آنتن) به یکی از ردیف‌های PMO؛ در غیر این صورت None."""
    vendor = (vendor or "").strip()
    antenna_type = (antenna_type or "").strip()
    if mapping_map is not None and antenna_type:
        mapped = mapping_map.get(antenna_type.casefold())
        if mapped is not None:
            if not mapped or mapped == PMO_SKIP_ROW:
                return None
            if mapped in PMO_ROW_ORDER:
                return mapped
    return _pmo_row_for_auto(vendor, antenna_type)


def _pmo_row_auto(
    vendor: str,
    antenna_type: str,
    item_code: str,
    item_type_map: dict[str, str],
    mapping_map: Dict[str, str] | None = None,
    item_mapping_map: Dict[tuple[str, str, str], str] | None = None,
) -> str | None:
    """ردیف PMO: اول نگاشت آیتم‌کد، بعد نوع آنتن، بعد خودکار."""
    row_key = _pmo_row_for_item(vendor, antenna_type, item_code, item_mapping_map)
    if row_key is not None:
        return row_key
    row_key = _pmo_row_for(vendor, antenna_type, mapping_map)
    if row_key is not None:
        return row_key
    code = (item_code or "").strip().lower()
    if code:
        mapped = item_type_map.get(code, "")
        if mapped:
            row_key = _pmo_row_for_item(vendor, mapped, item_code, item_mapping_map)
            if row_key is not None:
                return row_key
            return _pmo_row_for(vendor, mapped, mapping_map)
    return None


def _pmo_scope_map(db: Session) -> Dict[str, int]:
    """نگاشت {row_type: scope_qty} از جدول دستیِ PmoScope (SCOPE/کل تعمیری)."""
    result: Dict[str, int] = {}
    if not hasattr(models, "PmoScope"):
        return result
    for row in db.query(models.PmoScope).all():
        key = (row.row_type or "").strip()
        if key:
            result[key] = int(row.scope_qty or 0)
    return result


def _pmo_mapping_map(db: Session) -> Dict[str, str]:
    """نگاشت {antenna_type.casefold(): pmo_row_type} از جدول PmoAntennaMapping."""
    result: Dict[str, str] = {}
    if not hasattr(models, "PmoAntennaMapping"):
        return result
    for row in db.query(models.PmoAntennaMapping).all():
        atype = (row.antenna_type or "").strip()
        if atype:
            result[atype.casefold()] = (row.pmo_row_type or "").strip()
    return result


def _pmo_item_mapping_key(vendor: str, antenna_type: str, item_code: str) -> tuple[str, str, str]:
    return (
        (vendor or "").strip().casefold(),
        (antenna_type or "").strip().casefold(),
        (item_code or "").strip().casefold(),
    )


def _pmo_item_mapping_map(db: Session) -> Dict[tuple[str, str, str], str]:
    """نگاشت {(vendor, antenna_type, item_code): pmo_row_type}."""
    result: Dict[tuple[str, str, str], str] = {}
    if not hasattr(models, "PmoItemMapping"):
        return result
    for row in db.query(models.PmoItemMapping).all():
        code = (row.item_code or "").strip()
        if not code:
            continue
        key = _pmo_item_mapping_key(row.vendor, row.antenna_type, code)
        result[key] = (row.pmo_row_type or "").strip()
    return result


def _pmo_row_for_item(
    vendor: str,
    antenna_type: str,
    item_code: str,
    item_mapping_map: Dict[tuple[str, str, str], str] | None,
) -> str | None:
    code = (item_code or "").strip()
    if not code or not item_mapping_map:
        return None
    mapped = item_mapping_map.get(_pmo_item_mapping_key(vendor, antenna_type, code))
    if mapped is None:
        return None
    if not mapped or mapped == PMO_SKIP_ROW:
        return None
    if mapped in PMO_ROW_ORDER:
        return mapped
    return None


def _collect_pmo_item_triples(db: Session) -> List[Dict[str, str]]:
    """Distinct vendor + antenna_type + item_code seen in the system."""
    seen: set[tuple[str, str, str]] = set()
    items: List[Dict[str, str]] = []

    def add(vendor: str, antenna_type: str, item_code: str) -> None:
        vendor = (vendor or "").strip()
        antenna_type = (antenna_type or "").strip()
        item_code = (item_code or "").strip()
        if not item_code:
            return
        key = _pmo_item_mapping_key(vendor, antenna_type, item_code)
        if key in seen:
            return
        seen.add(key)
        items.append(
            {
                "vendor": vendor,
                "antenna_type": antenna_type,
                "item_code": item_code,
            }
        )

    if hasattr(models, "ItemCodeAssignment"):
        for row in db.query(models.ItemCodeAssignment).all():
            add(row.manufacturer, row.antenna_type, row.item_code)
    if hasattr(models, "Repair"):
        for row in db.query(
            models.Repair.constructor,
            models.Repair.antenna_type,
            models.Repair.item_code,
        ).distinct():
            add(row.constructor, row.antenna_type, row.item_code)
    if hasattr(models, "NoSerialAntenna"):
        for row in db.query(
            models.NoSerialAntenna.vendor,
            models.NoSerialAntenna.antenna_type,
            models.NoSerialAntenna.item_code,
        ).distinct():
            add(row.vendor, row.antenna_type, row.item_code)

    items.sort(
        key=lambda row: (
            row["item_code"].casefold(),
            row["vendor"].casefold(),
            row["antenna_type"].casefold(),
        )
    )
    return items


def _build_pmo_item_mappings_payload(db: Session) -> Dict[str, Any]:
    manual = _pmo_item_mapping_map(db)
    manual_rows = {
        (
            (row.vendor or "").strip(),
            (row.antenna_type or "").strip(),
            (row.item_code or "").strip(),
        ): (row.pmo_row_type or "").strip()
        for row in db.query(models.PmoItemMapping).all()
        if (row.item_code or "").strip()
    } if hasattr(models, "PmoItemMapping") else {}
    item_type_map = _item_code_to_antenna_type_map(db)
    antenna_mapping = _pmo_mapping_map(db)
    items: List[Dict[str, Any]] = []
    for triple in _collect_pmo_item_triples(db):
        vendor = triple["vendor"]
        atype = triple["antenna_type"]
        code = triple["item_code"]
        manual_val = manual_rows.get((vendor, atype, code), "")
        auto_val = (
            _pmo_row_auto(vendor, atype, code, item_type_map, antenna_mapping, None)
            or ""
        )
        effective = (
            _pmo_row_auto(
                vendor, atype, code, item_type_map, antenna_mapping, manual
            )
            or ""
        )
        items.append(
            {
                "vendor": vendor,
                "antenna_type": atype,
                "item_code": code,
                "pmo_row_type": manual_val,
                "auto_row_type": auto_val,
                "effective_row_type": effective,
            }
        )
    row_types = [
        {"code": code, "label": PMO_ROW_LABELS.get(code, code)}
        for code in PMO_ROW_ORDER
    ]
    return {
        "row_types": row_types,
        "skip_value": PMO_SKIP_ROW,
        "items": items,
    }


def _collect_pmo_antenna_types(db: Session) -> List[str]:
    """همهٔ انواع آنتن دیده‌شده در سیستم (مرتب‌شده)."""
    found: set[str] = set()
    if hasattr(models, "AntennaType"):
        for row in db.query(models.AntennaType.antenna_type).all():
            if row.antenna_type and row.antenna_type.strip():
                found.add(row.antenna_type.strip())
    if hasattr(models, "Repair"):
        for (atype,) in db.query(models.Repair.antenna_type).distinct():
            if atype and str(atype).strip():
                found.add(str(atype).strip())
    if hasattr(models, "NoSerialAntenna"):
        for (atype,) in db.query(models.NoSerialAntenna.antenna_type).distinct():
            if atype and str(atype).strip():
                found.add(str(atype).strip())
    if hasattr(models, "ItemCodeAssignment"):
        for (atype,) in db.query(models.ItemCodeAssignment.antenna_type).distinct():
            if atype and str(atype).strip():
                found.add(str(atype).strip())
    return sorted(found, key=lambda x: x.casefold())


def _pmo_repair_serial_key(repair_no: int | None, antenna_serial_no: str | None) -> str:
    """کلید یکتا برای شمارش PMO؛ سریال خالی → یک رکورد = یک واحد."""
    key = repair_serial_key(antenna_serial_no)
    if key:
        return key
    return f"__repair_{repair_no or 0}__"


def _live_sap_counts(db: Session) -> Dict[str, int]:
    """تعداد سریال یکتا با وضعیت خروج در repairs، به‌ازای هر سطر PMO (2x/3x/...)."""
    return {k: v["sap"] for k, v in _live_repair_status_cols(db).items()}


def _live_repair_status_cols(db: Session) -> Dict[str, Dict[str, int]]:
    """sap/wait/tot زنده از repairs: خروج و تایید، تفکیک‌شده روی سطرهای PMO."""
    sap_serials: Dict[str, set[str]] = defaultdict(set)
    wait_serials: Dict[str, set[str]] = defaultdict(set)
    item_type_map = _item_code_to_antenna_type_map(db)
    mapping_map = _pmo_mapping_map(db)
    item_mapping_map = _pmo_item_mapping_map(db)

    for rep in db.query(
        models.Repair.repair_no,
        models.Repair.antenna_serial_no,
        models.Repair.constructor,
        models.Repair.antenna_type,
        models.Repair.item_code,
        models.Repair.antenna_status,
    ).all():
        row_key = _pmo_row_auto(
            rep.constructor,
            rep.antenna_type,
            rep.item_code,
            item_type_map,
            mapping_map,
            item_mapping_map,
        )
        if row_key is None:
            continue
        serial = _pmo_repair_serial_key(rep.repair_no, rep.antenna_serial_no)
        status = rep.antenna_status or ""
        if _is_exit_status(status):
            sap_serials[row_key].add(serial)
        elif status == "تایید":
            wait_serials[row_key].add(serial)

    out: Dict[str, Dict[str, int]] = {}
    for name in PMO_ROW_ORDER:
        sap = len(sap_serials.get(name, set()))
        wait = len(wait_serials.get(name, set()) - sap_serials.get(name, set()))
        out[name] = {"sap": sap, "wait": wait, "tot": sap + wait}
    return out


def _overlay_live_sap(db: Session, rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """ستون‌های خروج/تایید/کل تعمیری را همیشه از جدول repairs زنده جایگزین می‌کند."""
    live = _live_repair_status_cols(db)
    out: List[Dict[str, Any]] = []
    for row in rows:
        merged = dict(row)
        row_type = merged.get("type")
        if row_type in live:
            merged["sap"] = live[row_type]["sap"]
            merged["wait"] = live[row_type]["wait"]
            merged["tot"] = live[row_type]["tot"]
        out.append(merged)
    return out


def _build_pmo_report(db: Session) -> List[Dict[str, Any]]:
    """جدول گزارش پروژه PMO، زنده از دیتابیس.

    برای هر کلاس آنتن:
      • dep  = جمع موجودی «مستهلک» از NoSerialAntenna
      • p1   = جمع «ورودی جدید از استان» از NoSerialAntenna (دسته «استان»)
      • p2   = همان موجودی مستهلک (اولویت ۲)
      • sap  = تعداد سریال یکتا با وضعیت «خروج»
      • wait = تعداد سریال یکتا با وضعیت «تایید» (بدون شمارش مجدد سریال‌های خروج‌خورده)
      • tot  = کل تعمیری = sap + wait (هر سریال حداکثر یک بار)
    """
    rows: Dict[str, Dict[str, int]] = {
        name: {"dep": 0, "p1": 0, "p2": 0, "tot": 0, "sap": 0, "wait": 0}
        for name in PMO_ROW_ORDER
    }
    sap_serials: Dict[str, set[str]] = defaultdict(set)
    wait_serials: Dict[str, set[str]] = defaultdict(set)
    item_type_map = _item_code_to_antenna_type_map(db)
    mapping_map = _pmo_mapping_map(db)
    item_mapping_map = _pmo_item_mapping_map(db)

    if hasattr(models, "NoSerialAntenna"):
        for r in db.query(
            models.NoSerialAntenna.vendor,
            models.NoSerialAntenna.antenna_type,
            models.NoSerialAntenna.item_code,
            models.NoSerialAntenna.quantity,
            models.NoSerialAntenna.remaining_quantity,
            models.NoSerialAntenna.category,
        ).all():
            row_key = _pmo_row_auto(
                r.vendor,
                r.antenna_type,
                r.item_code,
                item_type_map,
                mapping_map,
                item_mapping_map,
            )
            if row_key is None:
                continue
            qty = remaining_of(r)
            category = (r.category or "").strip()
            if category == CATEGORY_DEPRECIATED:
                rows[row_key]["dep"] += qty
            elif category == "استان":
                rows[row_key]["p1"] += qty

    for rep in db.query(
        models.Repair.repair_no,
        models.Repair.antenna_serial_no,
        models.Repair.constructor,
        models.Repair.antenna_type,
        models.Repair.item_code,
        models.Repair.antenna_status,
    ).all():
        row_key = _pmo_row_auto(
            rep.constructor,
            rep.antenna_type,
            rep.item_code,
            item_type_map,
            mapping_map,
            item_mapping_map,
        )
        if row_key is None:
            continue
        serial = _pmo_repair_serial_key(rep.repair_no, rep.antenna_serial_no)
        status = rep.antenna_status or ""
        if _is_exit_status(status):
            sap_serials[row_key].add(serial)
        elif status == "تایید":
            wait_serials[row_key].add(serial)

    result: List[Dict[str, Any]] = []
    for name in PMO_ROW_ORDER:
        data = rows[name]
        data["p2"] = data["dep"]
        sap = len(sap_serials[name])
        wait = len(wait_serials[name] - sap_serials[name])
        data["sap"] = sap
        data["wait"] = wait
        data["tot"] = sap + wait
        result.append({"type": name, **data})
    return result


@router.get("/summary")
async def dashboard_summary(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """Fresh aggregated stats for all dashboard widgets."""
    response.headers["Cache-Control"] = "no-store"

    total_users = (
        db.query(models.User)
        .filter(
            models.User.is_active.is_(True),
            models.User.user_type == "regularUser",
        )
        .count()
    )

    inventory = (
        db.query(models.ArsenalTable)
        .order_by(models.ArsenalTable.fragment_name)
        .all()
    )

    entry_totals: dict[str, int] = defaultdict(int)
    for entry in db.query(models.SupplyFragmentEntry).all():
        if not entry.fragment_part_no or not entry.arsenal_name:
            continue
        key = f"{entry.fragment_part_no}::{entry.arsenal_name}"
        entry_totals[key] += int(entry.fragment_entry_qty or 0)

    inventory_preview = [
        {
            "fragment_name": item.fragment_name,
            "fragment_qty": int(item.fragment_qty or 0),
            "arsenal_name": item.arsenal_name,
        }
        for item in sorted(
            (i for i in inventory if i.arsenal_name != ARSENAL_EQUIPMENT),
            key=lambda row: int(row.fragment_qty or 0),
            reverse=True,
        )[:5]
    ]

    inventory_chart: List[Dict[str, Any]] = []
    for item in sorted(
        (i for i in inventory if i.arsenal_name != ARSENAL_EQUIPMENT),
        key=lambda row: int(row.fragment_qty or 0),
        reverse=True,
    )[:15]:
        key = f"{item.fragment_part_no}::{item.arsenal_name}"
        total_in = entry_totals.get(key, int(item.fragment_qty or 0))
        stock = int(item.fragment_qty or 0)
        name = item.fragment_name or ""
        inventory_chart.append(
            {
                "key": key,
                "fullName": name,
                "label": name,
                "stock": stock,
                "totalIn": total_in,
                "consumed": max(0, total_in - stock),
            }
        )

    pending_approval = (
        db.query(models.Repair)
        .filter(models.Repair.antenna_status == "در_انتظار_تایید")
        .count()
    )

    queue_repairs = (
        db.query(models.Repair)
        .filter(models.Repair.antenna_status == "نامعلوم")
        .count()
    )

    queue_rows = (
        db.query(
            models.Repair.constructor,
            models.Repair.antenna_type,
            func.count(models.Repair.repair_no).label("cnt"),
        )
        .filter(models.Repair.antenna_status == "نامعلوم")
        .group_by(models.Repair.constructor, models.Repair.antenna_type)
        .order_by(func.count(models.Repair.repair_no).desc())
        .limit(5)
        .all()
    )

    queue_preview = [
        {
            "vendor": (row.constructor or "-").strip() or "-",
            "antenna_type": (row.antenna_type or "-").strip() or "-",
            "count": int(row.cnt or 0),
        }
        for row in queue_rows
    ]

    status_counts: dict[str, int] = defaultdict(int)
    for repair in db.query(models.Repair.antenna_status).all():
        status = repair.antenna_status or "نامعلوم"
        status_counts[status] += 1

    repair_status_chart = [
        {
            "status": status,
            "name": _status_label(status),
            "value": count,
        }
        for status, count in status_counts.items()
    ]

    item_code_rows = (
        db.query(
            models.Repair.item_code,
            func.count(models.Repair.repair_no).label("cnt"),
        )
        .filter(
            models.Repair.item_code.isnot(None),
            models.Repair.item_code != "",
        )
        .group_by(models.Repair.item_code)
        .order_by(func.count(models.Repair.repair_no).desc())
        .all()
    )

    item_codes_preview = [
        {
            "code": (row.item_code or "").strip(),
            "count": int(row.cnt or 0),
        }
        for row in item_code_rows[:5]
        if (row.item_code or "").strip()
    ]

    pmo_rows = _pmo_snapshot_rows(db)
    antenna_type_x_bars = bars_from_snapshot(pmo_rows)
    pmo_cards = snapshot_card_totals(pmo_rows)

    return {
        "total_users": total_users,
        "inventory_items": len(inventory),
        "inventory_preview": inventory_preview,
        "pending_approval": pending_approval,
        "queue_repairs": queue_repairs,
        "queue_preview": queue_preview,
        "inventory_chart": inventory_chart,
        "repair_status_chart": repair_status_chart,
        "item_code_count": len(item_code_rows),
        "item_codes_preview": item_codes_preview,
        "item_code_bars": _build_item_code_bars(db),
        "antenna_type_x_bars": antenna_type_x_bars,
        "antenna_center_inventory_bars": center_inventory_from_snapshot(pmo_rows),
        "sap_pie_bars": antenna_type_x_bars,
        "antenna_targets": _build_antenna_targets_map(db),
        "top_repairers": _build_top_repairers(db),
        "vendor_top_units": _build_vendor_top_units(db),
        "province_incoming_total": pmo_cards["province_incoming_total"],
        "center_antennas": pmo_cards["center_antennas"],
        "pmo_source": "snapshot+live_sap",
    }


def _pmo_snapshot_rows(db: Session) -> List[Dict[str, Any]]:
    """اعداد PMO ذخیره‌شده؛ فیلد sap همیشه از repairs زنده جایگزین می‌شود."""
    if snapshot_has_rows(db):
        rows = load_snapshot(db, PMO_ROW_ORDER)
    else:
        rows = seed_if_empty(db, _build_pmo_report(db), PMO_ROW_ORDER)
    return _overlay_live_sap(db, rows)


def _pmo_report_payload(db: Session) -> Dict[str, Any]:
    return {
        "rows": _pmo_snapshot_rows(db),
        "has_overrides": True,
        "source": "pmo_snapshot+live_sap",
        "sap_source": "repairs_exit",
    }


@router.get("/pmo-report")
async def dashboard_pmo_report(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """گزارش PMO: بقیه فیلدها از snapshot، Done(SAP) زنده از repairs با وضعیت خروج."""
    response.headers["Cache-Control"] = "no-store"
    return _pmo_report_payload(db)


@router.put("/pmo-display")
async def dashboard_pmo_display_save(
    response: Response,
    payload: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_admin_user),
) -> Dict[str, Any]:
    """ذخیرهٔ موجودی/اولویت‌ها؛ خروج/تایید/کل همیشه از repairs محاسبه و ذخیره می‌شود."""
    response.headers["Cache-Control"] = "no-store"
    rows = payload.get("rows")
    if not isinstance(rows, list):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="rows must be a list",
        )
    # ignore client sap/wait/tot — always take live exit/approve counts from repairs
    rows = _overlay_live_sap(db, rows)
    save_snapshot(db, rows, PMO_ROW_ORDER)
    clear_pmo_display_overrides()
    return {
        "rows": _pmo_snapshot_rows(db),
        "has_overrides": True,
        "source": "pmo_snapshot+live_sap",
        "sap_source": "repairs_exit",
    }


@router.delete("/pmo-display")
async def dashboard_pmo_display_reset(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_admin_user),
) -> Dict[str, Any]:
    """بازنشانی override نمایش؛ Done(SAP) همچنان از repairs زنده برمی‌گردد."""
    response.headers["Cache-Control"] = "no-store"
    clear_pmo_display_overrides()
    return {
        "rows": _pmo_snapshot_rows(db),
        "has_overrides": True,
        "source": "pmo_snapshot+live_sap",
        "sap_source": "repairs_exit",
    }


@router.put("/pmo-scope")
async def dashboard_pmo_scope_save(
    response: Response,
    payload: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_admin_user),
) -> Dict[str, Any]:
    """ذخیرهٔ عدد SCOPE (کل تعمیری) به‌صورت دستی برای ردیف‌های PMO.

    بدنه: {"scopes": {"x-pole": 132, "2x-pole": 373, ...}}
    """
    response.headers["Cache-Control"] = "no-store"
    scopes = payload.get("scopes")
    if not isinstance(scopes, dict):
        scopes = {
            str(row.get("type")): row.get("scope", row.get("tot"))
            for row in (payload.get("rows") or [])
            if isinstance(row, dict) and row.get("type")
        }

    valid_types = set(PMO_ROW_ORDER)
    existing = {
        (row.row_type or "").strip(): row
        for row in db.query(models.PmoScope).all()
    }
    for row_type, raw_qty in scopes.items():
        row_type = (str(row_type) or "").strip()
        if row_type not in valid_types:
            continue
        try:
            qty = max(0, int(raw_qty or 0))
        except (TypeError, ValueError):
            qty = 0
        record = existing.get(row_type)
        if record is None:
            db.add(models.PmoScope(row_type=row_type, scope_qty=qty))
        else:
            record.scope_qty = qty
    db.commit()
    payload = _pmo_report_payload(db)
    return {"rows": payload["rows"]}


def _build_pmo_mappings_payload(db: Session) -> Dict[str, Any]:
    manual = _pmo_mapping_map(db)
    manual_by_type = {
        (row.antenna_type or "").strip(): (row.pmo_row_type or "").strip()
        for row in db.query(models.PmoAntennaMapping).all()
        if (row.antenna_type or "").strip()
    }
    items: List[Dict[str, Any]] = []
    for atype in _collect_pmo_antenna_types(db):
        manual_val = manual_by_type.get(atype, "")
        auto_val = _pmo_row_for_auto("", atype) or ""
        effective = _pmo_row_for("", atype, manual) or ""
        items.append(
            {
                "antenna_type": atype,
                "pmo_row_type": manual_val,
                "auto_row_type": auto_val,
                "effective_row_type": effective,
            }
        )
    row_types = [
        {"code": code, "label": PMO_ROW_LABELS.get(code, code)}
        for code in PMO_ROW_ORDER
    ]
    return {
        "row_types": row_types,
        "skip_value": PMO_SKIP_ROW,
        "items": items,
    }


@router.get("/pmo-mappings")
@router.get("/pmo_mappings")
async def dashboard_pmo_mappings_get(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """لیست انواع آنتن و نگاشت دستی/خودکار به ستون‌های PMO."""
    response.headers["Cache-Control"] = "no-store"
    return _build_pmo_mappings_payload(db)


@router.put("/pmo-mappings")
@router.put("/pmo_mappings")
async def dashboard_pmo_mappings_save(
    response: Response,
    payload: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_admin_user),
) -> Dict[str, Any]:
    """ذخیرهٔ نگاشت نوع آنتن → ستون PMO.

    بدنه: {"mappings": {"2X-Kathrein": "2x-pole", "MQ4": "__skip__", ...}}
    مقدار خالی = حذف نگاشت دستی (بازگشت به خودکار).
    """
    response.headers["Cache-Control"] = "no-store"
    mappings = payload.get("mappings")
    if not isinstance(mappings, dict):
        mappings = {
            str(row.get("antenna_type")): row.get("pmo_row_type", "")
            for row in (payload.get("items") or [])
            if isinstance(row, dict) and row.get("antenna_type")
        }

    valid_rows = set(PMO_ROW_ORDER) | {PMO_SKIP_ROW, ""}
    existing = {
        (row.antenna_type or "").strip(): row
        for row in db.query(models.PmoAntennaMapping).all()
    }
    now = jdatetime.datetime.now().strftime("%Y/%m/%d %H:%M:%S")

    for raw_type, raw_row in mappings.items():
        atype = (str(raw_type) or "").strip()
        if not atype:
            continue
        row_type = (str(raw_row) if raw_row is not None else "").strip()
        if row_type and row_type not in valid_rows:
            continue
        record = existing.get(atype)
        if not row_type:
            if record is not None:
                db.delete(record)
            continue
        if record is None:
            db.add(
                models.PmoAntennaMapping(
                    antenna_type=atype,
                    pmo_row_type=row_type,
                    updated_date=now,
                )
            )
        else:
            record.pmo_row_type = row_type
            record.updated_date = now

    db.commit()
    pmo_payload = _pmo_report_payload(db)
    return {
        "rows": pmo_payload["rows"],
        "has_overrides": pmo_payload["has_overrides"],
        **_build_pmo_mappings_payload(db),
    }


@router.get("/pmo-item-mappings")
@router.get("/pmo_item_mappings")
async def dashboard_pmo_item_mappings_get(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """لیست آیتم‌کدها با نگاشت دستی/خودکار به بسته PMO."""
    response.headers["Cache-Control"] = "no-store"
    return _build_pmo_item_mappings_payload(db)


@router.put("/pmo-item-mappings")
@router.put("/pmo_item_mappings")
async def dashboard_pmo_item_mappings_save(
    response: Response,
    payload: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_admin_user),
) -> Dict[str, Any]:
    """ذخیرهٔ نگاشت (وندور + نوع آنتن + آیتم‌کد) → بسته PMO."""
    response.headers["Cache-Control"] = "no-store"
    if not hasattr(models, "PmoItemMapping"):
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="PmoItemMapping table is not available",
        )

    items = payload.get("items")
    if not isinstance(items, list):
        items = []
    rebuild = bool(payload.get("rebuild_snapshot"))

    valid_rows = set(PMO_ROW_ORDER) | {PMO_SKIP_ROW, ""}
    existing = {
        (
            (row.vendor or "").strip(),
            (row.antenna_type or "").strip(),
            (row.item_code or "").strip(),
        ): row
        for row in db.query(models.PmoItemMapping).all()
        if (row.item_code or "").strip()
    }
    now = jdatetime.datetime.now().strftime("%Y/%m/%d %H:%M:%S")
    touched: set[tuple[str, str, str]] = set()

    for raw in items:
        if not isinstance(raw, dict):
            continue
        vendor = (str(raw.get("vendor") or "")).strip()
        atype = (str(raw.get("antenna_type") or "")).strip()
        code = (str(raw.get("item_code") or "")).strip()
        if not code:
            continue
        row_type = (str(raw.get("pmo_row_type") if raw.get("pmo_row_type") is not None else "")).strip()
        if row_type and row_type not in valid_rows:
            continue
        key = (vendor, atype, code)
        touched.add(key)
        record = existing.get(key)
        if not row_type:
            if record is not None:
                db.delete(record)
            continue
        if record is None:
            db.add(
                models.PmoItemMapping(
                    vendor=vendor,
                    antenna_type=atype,
                    item_code=code,
                    pmo_row_type=row_type,
                    updated_date=now,
                )
            )
        else:
            record.pmo_row_type = row_type
            record.updated_date = now

    db.commit()

    if rebuild:
        live_rows = _build_pmo_report(db)
        save_snapshot(db, live_rows, PMO_ROW_ORDER)

    pmo_payload = _pmo_report_payload(db)
    return {
        "rows": pmo_payload["rows"],
        "has_overrides": pmo_payload["has_overrides"],
        "rebuilt": rebuild,
        **_build_pmo_item_mappings_payload(db),
    }


@router.post("/pmo-report/rebuild-from-db")
async def dashboard_pmo_report_rebuild_from_db(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_admin_user),
) -> Dict[str, Any]:
    """بازمحاسبه snapshot PMO از دیتابیس با نگاشت‌های فعلی (نوع + آیتم‌کد)."""
    response.headers["Cache-Control"] = "no-store"
    live_rows = _build_pmo_report(db)
    save_snapshot(db, live_rows, PMO_ROW_ORDER)
    return {
        "message": "PMO snapshot rebuilt from database",
        **_pmo_report_payload(db),
    }


@router.get("/province-inventory")
async def dashboard_province_inventory(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """موجودی «ورودی استان» مرکز (زنده از دیتابیس) برای پنل آمار موجودی."""
    response.headers["Cache-Control"] = "no-store"
    return _build_province_inventory(db)


@router.get("/depreciated-inventory")
async def dashboard_depreciated_inventory(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """موجودی مستهلک مرکز (زنده از دیتابیس) برای پنل «آمار موجودی در مرکز آنتن»."""
    response.headers["Cache-Control"] = "no-store"
    return _build_depreciated_inventory(db)


@router.get("/top-repairers")
async def dashboard_top_repairers(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """Per-user top 5 repair items for dashboard carousel."""
    response.headers["Cache-Control"] = "no-store"
    return {"repairers": _build_top_repairers(db)}


@router.get("/item-codes")
async def dashboard_item_codes(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """Distinct item codes with antenna counts."""
    response.headers["Cache-Control"] = "no-store"

    rows = (
        db.query(
            models.Repair.item_code,
            func.count(models.Repair.repair_no).label("cnt"),
        )
        .filter(
            models.Repair.item_code.isnot(None),
            models.Repair.item_code != "",
        )
        .group_by(models.Repair.item_code)
        .order_by(models.Repair.item_code)
        .all()
    )

    codes = [
        {
            "code": (row.item_code or "").strip(),
            "count": int(row.cnt or 0),
        }
        for row in rows
        if (row.item_code or "").strip()
    ]

    return {"codes": codes}


@router.get("/item-code-status")
async def dashboard_item_code_status(
    response: Response,
    code: str = Query(..., min_length=1),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """Status breakdown for antennas with a given item code."""
    response.headers["Cache-Control"] = "no-store"

    code_value = code.strip()
    status_counts: dict[str, int] = defaultdict(int)

    for repair in (
        db.query(models.Repair)
        .filter(models.Repair.item_code == code_value)
        .all()
    ):
        status = repair.antenna_status or "نامعلوم"
        status_counts[status] += 1

    total = sum(status_counts.values())

    return {
        "code": code_value,
        "total": total,
        "chart": _build_item_code_chart(status_counts),
    }


@router.get("/item-codes-overview")
async def dashboard_item_codes_overview(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """All item codes with status charts, antenna type, and vendor."""
    response.headers["Cache-Control"] = "no-store"

    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "status_counts": defaultdict(int),
            "type_counts": defaultdict(int),
            "vendor_counts": defaultdict(int),
            "count": 0,
        }
    )

    for repair in (
        db.query(
            models.Repair.item_code,
            models.Repair.antenna_status,
            models.Repair.antenna_type,
            models.Repair.constructor,
        )
        .filter(
            models.Repair.item_code.isnot(None),
            models.Repair.item_code != "",
        )
        .all()
    ):
        code = (repair.item_code or "").strip()
        if not code:
            continue
        bucket = grouped[code]
        bucket["count"] += 1
        bucket["status_counts"][repair.antenna_status or "نامعلوم"] += 1
        antenna_type = (repair.antenna_type or "").strip()
        if antenna_type:
            bucket["type_counts"][antenna_type] += 1
        vendor = (repair.constructor or "").strip()
        if vendor:
            bucket["vendor_counts"][vendor] += 1

    items: List[Dict[str, Any]] = []
    for code, data in sorted(grouped.items(), key=lambda row: -row[1]["count"]):
        total = int(data["count"])
        items.append(
            {
                "code": code,
                "count": total,
                "total": total,
                "antenna_type": _mode_label(data["type_counts"]),
                "vendor": _mode_label(data["vendor_counts"]),
                "chart": _build_item_code_chart(data["status_counts"]),
            }
        )

    return {"items": items}


def _has_exit_date(value: Any) -> bool:
    return bool(str(value or "").strip())


def _item_code_report_query(
    db: Session,
    *,
    item_code: str | None = None,
    vendor: str | None = None,
    antenna_type: str | None = None,
    status: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    date_field: str = "exit",
    only_exited: bool = False,
    exact_item_code: str | None = None,
):
    query = db.query(models.Repair)

    code_filter = (item_code or "").strip()
    vendor_filter = (vendor or "").strip()
    type_filter = (antenna_type or "").strip()
    status_filter = (status or "").strip()
    exact_code = (exact_item_code or "").strip()

    if exact_code:
        query = query.filter(models.Repair.item_code == exact_code)
    elif code_filter:
        query = query.filter(models.Repair.item_code.like(f"%{code_filter}%"))
    if vendor_filter:
        query = query.filter(models.Repair.constructor.like(f"%{vendor_filter}%"))
    if type_filter:
        query = query.filter(models.Repair.antenna_type.like(f"%{type_filter}%"))
    if status_filter:
        query = query.filter(models.Repair.antenna_status == status_filter)
    if only_exited:
        query = query.filter(
            or_(
                models.Repair.antenna_status.in_(list(EXIT_STATUSES)),
                models.Repair.antenna_status.like("خروج%"),
                and_(
                    models.Repair.exit_date.isnot(None),
                    models.Repair.exit_date != "",
                ),
            )
        )

    field = (date_field or "exit").strip().lower()
    date_column = models.Repair.exit_date if field == "exit" else models.Repair.entry_date
    if start_date or end_date:
        query = apply_jalali_date_range_filter(query, date_column, start_date, end_date)
    return query


def _item_code_report_filter_rows(**filters: Any) -> List[List[str]]:
    date_field = (filters.get("date_field") or "exit").strip().lower()
    date_label = "تاریخ خروج" if date_field == "exit" else "تاریخ ورود"
    status = (filters.get("status") or "").strip()
    status_label = _status_label(status) if status else "همه وضعیت‌ها"
    rows = [
        ["آیتم کد", (filters.get("item_code") or "").strip() or "همه"],
        ["وندور", (filters.get("vendor") or "").strip() or "همه"],
        ["نوع آنتن", (filters.get("antenna_type") or "").strip() or "همه"],
        ["وضعیت", status_label],
        ["از تاریخ", (filters.get("start_date") or "").strip() or "—"],
        ["تا تاریخ", (filters.get("end_date") or "").strip() or "—"],
        ["فیلتر تاریخ روی", date_label],
        ["فقط دارای خروج", "بله" if filters.get("only_exited") else "خیر"],
        ["تاریخ گزارش", jdatetime.datetime.now().strftime("%Y/%m/%d %H:%M")],
    ]
    return rows


def _collect_item_code_serial_rows(db: Session, **filters: Any) -> List[Dict[str, Any]]:
    query = _item_code_report_query(db, **filters).order_by(
        models.Repair.item_code.asc(),
        models.Repair.repair_no.desc(),
    )
    rows: List[Dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for repair in query.all():
        code = (repair.item_code or "").strip()
        if not code:
            continue
        serial = (repair.antenna_serial_no or "").strip()
        key = (code, repair_serial_key(serial) or f"#{repair.repair_no}")
        if key in seen:
            continue
        seen.add(key)
        status = repair.antenna_status or "نامعلوم"
        rows.append(
            {
                "item_code": code,
                "repair_no": repair.repair_no,
                "serial": serial,
                "status": status,
                "status_label": _status_label(status),
                "vendor": repair.constructor or "",
                "antenna_type": repair.antenna_type or "",
                "entry_date": repair.entry_date or "",
                "exit_date": repair.exit_date or "",
            }
        )
    return rows


def _style_excel_header_row(ws, row_idx: int, col_count: int) -> None:
    fill = PatternFill("solid", fgColor="E7D6A8")
    font = Font(bold=True)
    align = Alignment(horizontal="center", vertical="center")
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align


def _autosize_worksheet_columns(ws, min_width: int = 10, max_width: int = 42) -> None:
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        width = min_width
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 2, max_width))
        ws.column_dimensions[letter].width = width


def _build_item_code_export_workbook(db: Session, **filters: Any) -> BytesIO:
    report = _build_item_code_status_report(db, **filters)
    serial_rows = _collect_item_code_serial_rows(db, **filters)
    summary = report.get("summary") or {}
    status_columns: List[Dict[str, str]] = report.get("status_columns") or []
    items: List[Dict[str, Any]] = report.get("items") or []

    wb = Workbook()

    ws_filters = wb.active
    ws_filters.title = "فیلترها"
    ws_filters.sheet_view.rightToLeft = True
    ws_filters.append(["عنوان", "مقدار"])
    for row in _item_code_report_filter_rows(**filters):
        ws_filters.append(row)
    ws_filters.append(["تعداد آیتم‌کد", summary.get("item_codes", 0)])
    ws_filters.append(["کل سریال‌ها", summary.get("serials", 0)])
    ws_filters.append(["سریال در وضعیت خروج", summary.get("exited", 0)])
    ws_filters.append(["دارای تاریخ خروج", summary.get("with_exit_date", 0)])
    ws_filters.append(["در مرکز", summary.get("in_center", 0)])
    _style_excel_header_row(ws_filters, 1, 2)
    _autosize_worksheet_columns(ws_filters)

    ws_summary = wb.create_sheet("خلاصه آیتم‌کد")
    ws_summary.sheet_view.rightToLeft = True
    summary_headers = [
        "آیتم کد",
        "وندور",
        "نوع آنتن",
        "کل سریال",
        "تعداد خروج",
        "دارای تاریخ خروج",
        "در مرکز",
    ] + [col["label"] for col in status_columns]
    ws_summary.append(summary_headers)
    for item in items:
        by_status = item.get("by_status") or {}
        ws_summary.append(
            [
                item.get("item_code", ""),
                item.get("vendor", ""),
                item.get("antenna_type", ""),
                item.get("serials", 0),
                item.get("exited", 0),
                item.get("with_exit_date", 0),
                item.get("in_center", 0),
            ]
            + [by_status.get(col["key"], 0) for col in status_columns]
        )
    _style_excel_header_row(ws_summary, 1, len(summary_headers))
    _autosize_worksheet_columns(ws_summary)

    ws_serials = wb.create_sheet("جزئیات سریال")
    ws_serials.sheet_view.rightToLeft = True
    serial_headers = [
        "آیتم کد",
        "شماره تعمیر",
        "سریال",
        "وضعیت",
        "وندور",
        "نوع آنتن",
        "تاریخ ورود",
        "تاریخ خروج",
    ]
    ws_serials.append(serial_headers)
    for row in serial_rows:
        ws_serials.append(
            [
                row["item_code"],
                row["repair_no"],
                row["serial"],
                row["status_label"],
                row["vendor"],
                row["antenna_type"],
                row["entry_date"],
                row["exit_date"],
            ]
        )
    _style_excel_header_row(ws_serials, 1, len(serial_headers))
    _autosize_worksheet_columns(ws_serials)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_item_code_status_report(
    db: Session,
    *,
    item_code: str | None = None,
    vendor: str | None = None,
    antenna_type: str | None = None,
    status: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    date_field: str = "exit",
    only_exited: bool = False,
) -> Dict[str, Any]:
    """Live serial counts per item code, grouped by antenna status."""
    query = _item_code_report_query(
        db,
        item_code=item_code,
        vendor=vendor,
        antenna_type=antenna_type,
        status=status,
        start_date=start_date,
        end_date=end_date,
        date_field=date_field,
        only_exited=only_exited,
    )

    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "status_counts": defaultdict(int),
            "type_counts": defaultdict(int),
            "vendor_counts": defaultdict(int),
            "serials": set(),
            "exited": 0,
            "with_exit_date": 0,
        }
    )
    seen: set[tuple[str, str]] = set()
    status_totals: dict[str, int] = defaultdict(int)
    missing_item_code = 0
    total_serials = 0
    total_exited = 0
    total_with_exit_date = 0

    for repair in query.all():
        code = (repair.item_code or "").strip()
        if not code:
            missing_item_code += 1
            continue
        serial = repair_serial_key(repair.antenna_serial_no) or f"#{repair.repair_no}"
        key = (code, serial)
        if key in seen:
            continue
        seen.add(key)

        row_status = (repair.antenna_status or "نامعلوم").strip() or "نامعلوم"
        exited = _is_exit_status(row_status)
        has_date = _has_exit_date(repair.exit_date)
        bucket = grouped[code]
        bucket["serials"].add(serial)
        bucket["status_counts"][row_status] += 1
        if exited:
            bucket["exited"] += 1
        if has_date:
            bucket["with_exit_date"] += 1
        antenna_type_value = (repair.antenna_type or "").strip()
        if antenna_type_value:
            bucket["type_counts"][antenna_type_value] += 1
        vendor_value = (repair.constructor or "").strip()
        if vendor_value:
            bucket["vendor_counts"][vendor_value] += 1

        status_totals[row_status] += 1
        total_serials += 1
        if exited:
            total_exited += 1
        if has_date:
            total_with_exit_date += 1

    used_statuses = set(status_totals)
    status_columns: List[Dict[str, str]] = []
    seen_status: set[str] = set()
    for key_status in ITEM_CODE_STATUS_ORDER:
        if key_status in used_statuses:
            status_columns.append(
                {"key": key_status, "label": _status_label(key_status)}
            )
            seen_status.add(key_status)
    for key_status in sorted(used_statuses):
        if key_status in seen_status:
            continue
        status_columns.append(
            {"key": key_status, "label": _status_label(key_status)}
        )

    items: List[Dict[str, Any]] = []
    for code, data in grouped.items():
        serial_count = len(data["serials"])
        exited = int(data["exited"])
        items.append(
            {
                "item_code": code,
                "vendor": _mode_label(data["vendor_counts"]),
                "antenna_type": _mode_label(data["type_counts"]),
                "serials": serial_count,
                "exited": exited,
                "with_exit_date": int(data["with_exit_date"]),
                "in_center": max(0, serial_count - exited),
                "by_status": {
                    key: int(count) for key, count in data["status_counts"].items()
                },
            }
        )

    items.sort(key=lambda row: (-int(row["exited"]), -int(row["serials"]), row["item_code"]))

    return {
        "summary": {
            "item_codes": len(items),
            "serials": total_serials,
            "exited": total_exited,
            "with_exit_date": total_with_exit_date,
            "in_center": max(0, total_serials - total_exited),
            "missing_item_code": missing_item_code,
            "by_status": [
                {
                    "status": col["key"],
                    "label": col["label"],
                    "count": int(status_totals.get(col["key"], 0)),
                }
                for col in status_columns
            ],
        },
        "status_columns": status_columns,
        "items": items,
    }


@router.get("/item-code-report")
async def dashboard_item_code_report(
    response: Response,
    item_code: Optional[str] = None,
    vendor: Optional[str] = None,
    antenna_type: Optional[str] = None,
    status: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    date_field: str = Query("exit"),
    only_exited: bool = Query(False),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """Search/reports window: item-code counts by serial status, live from repairs."""
    response.headers["Cache-Control"] = "no-store"
    return _build_item_code_status_report(
        db,
        item_code=item_code,
        vendor=vendor,
        antenna_type=antenna_type,
        status=status,
        start_date=start_date,
        end_date=end_date,
        date_field=date_field,
        only_exited=only_exited,
    )


@router.get("/item-code-report/serials")
async def dashboard_item_code_report_serials(
    response: Response,
    item_code: str = Query(..., min_length=1),
    vendor: Optional[str] = None,
    antenna_type: Optional[str] = None,
    status: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    date_field: str = Query("exit"),
    only_exited: bool = Query(False),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
) -> Dict[str, Any]:
    """Serial rows for one item code, same live filters as the report."""
    response.headers["Cache-Control"] = "no-store"
    code_value = item_code.strip()
    query = _item_code_report_query(
        db,
        item_code=None,
        vendor=vendor,
        antenna_type=antenna_type,
        status=status,
        start_date=start_date,
        end_date=end_date,
        date_field=date_field,
        only_exited=only_exited,
        exact_item_code=code_value,
    ).order_by(models.Repair.repair_no.desc())
    rows = []
    seen: set[str] = set()
    for repair in query.all():
        serial = (repair.antenna_serial_no or "").strip()
        key = repair_serial_key(serial) or f"#{repair.repair_no}"
        if key in seen:
            continue
        seen.add(key)
        rows.append(
            {
                "repair_no": repair.repair_no,
                "serial": serial,
                "status": repair.antenna_status or "نامعلوم",
                "status_label": _status_label(repair.antenna_status or "نامعلوم"),
                "vendor": repair.constructor or "",
                "antenna_type": repair.antenna_type or "",
                "entry_date": repair.entry_date or "",
                "exit_date": repair.exit_date or "",
            }
        )

    return {"item_code": code_value, "count": len(rows), "serials": rows}


@router.get("/item-code-report/export")
async def dashboard_item_code_report_export(
    item_code: Optional[str] = None,
    vendor: Optional[str] = None,
    antenna_type: Optional[str] = None,
    status: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    date_field: str = Query("exit"),
    only_exited: bool = Query(False),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """Excel export for item-code report: filters, summary, and serial details."""
    filters = {
        "item_code": item_code,
        "vendor": vendor,
        "antenna_type": antenna_type,
        "status": status,
        "start_date": start_date,
        "end_date": end_date,
        "date_field": date_field,
        "only_exited": only_exited,
    }
    report = _build_item_code_status_report(db, **filters)
    if not report.get("items"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="داده‌ای برای خروجی با این فیلتر وجود ندارد",
        )
    buf = _build_item_code_export_workbook(db, **filters)
    filename = f"item_code_report_{jdatetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
